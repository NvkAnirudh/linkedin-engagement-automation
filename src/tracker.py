"""Excel tracker for LinkedIn posts and comments."""

import json
import os
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Set
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from colorama import Fore, Style


class PostTracker:
    """Track LinkedIn posts and generated comments in Excel."""

    def __init__(self, excel_file: str = "data/tracked_posts.xlsx", cache_file: str = "data/seen_posts.json"):
        """Initialize post tracker.

        Args:
            excel_file: Path to Excel tracking file
            cache_file: Path to seen posts cache file
        """
        self.excel_file = excel_file
        self.cache_file = cache_file

        # Ensure data directory exists
        Path(excel_file).parent.mkdir(parents=True, exist_ok=True)

        # Initialize Excel file if it doesn't exist
        if not os.path.exists(excel_file):
            self._create_excel_file()

        # Load seen posts cache
        self.seen_posts = self._load_seen_posts()

    def _create_excel_file(self):
        """Create a new Excel file with proper headers."""
        print(f"{Fore.CYAN}Creating new Excel tracking file...{Style.RESET_ALL}")

        # Define columns
        columns = [
            'Post ID',
            'Profile Name',
            'Profile URL',
            'Post URL',
            'Post Date',
            'Scraped At',
            'Post Content',
            'Detected Tone',
            'Generated Comment',
            'Comment Style',
            'Reasoning',
            'Likes',
            'Comments',
            'Shares',
            'Posted Comment',
            'Posted At',
            'Notes'
        ]

        # Create DataFrame
        df = pd.DataFrame(columns=columns)

        # Save to Excel
        df.to_excel(self.excel_file, index=False, sheet_name='Posts')

        # Format the Excel file
        self._format_excel()

        print(f"{Fore.GREEN}Excel file created: {self.excel_file}{Style.RESET_ALL}")

    def _format_excel(self):
        """Apply formatting to Excel file."""
        try:
            workbook = load_workbook(self.excel_file)
            worksheet = workbook.active

            # Header formatting
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Set column widths
            column_widths = {
                'A': 15,  # Post ID
                'B': 20,  # Profile Name
                'C': 35,  # Profile URL
                'D': 35,  # Post URL
                'E': 15,  # Post Date
                'F': 15,  # Scraped At
                'G': 50,  # Post Content
                'H': 15,  # Detected Tone
                'I': 50,  # Generated Comment
                'J': 15,  # Comment Style
                'K': 40,  # Reasoning
                'L': 10,  # Likes
                'M': 10,  # Comments
                'N': 10,  # Shares
                'O': 10,  # Posted Comment
                'P': 15,  # Posted At
                'Q': 30,  # Notes
            }

            for col, width in column_widths.items():
                worksheet.column_dimensions[col].width = width

            # Enable text wrapping for content columns
            for row in worksheet.iter_rows(min_row=2):
                row[6].alignment = Alignment(wrap_text=True, vertical="top")  # Post Content
                row[8].alignment = Alignment(wrap_text=True, vertical="top")  # Generated Comment
                row[10].alignment = Alignment(wrap_text=True, vertical="top")  # Reasoning

            workbook.save(self.excel_file)

        except Exception as e:
            print(f"{Fore.YELLOW}Warning: Could not format Excel file: {str(e)}{Style.RESET_ALL}")

    def _load_seen_posts(self) -> Set[str]:
        """Load seen posts from cache file.

        Returns:
            Set of seen post IDs
        """
        if os.path.exists(self.cache_file):
            try:
                with open(self.cache_file, 'r') as f:
                    data = json.load(f)
                    return set(data.get('seen_posts', []))
            except Exception as e:
                print(f"{Fore.YELLOW}Warning: Could not load cache: {str(e)}{Style.RESET_ALL}")

        return set()

    def _save_seen_posts(self):
        """Save seen posts to cache file."""
        try:
            with open(self.cache_file, 'w') as f:
                json.dump({'seen_posts': list(self.seen_posts)}, f, indent=2)
        except Exception as e:
            print(f"{Fore.YELLOW}Warning: Could not save cache: {str(e)}{Style.RESET_ALL}")

    def is_post_seen(self, post_id: str) -> bool:
        """Check if a post has been seen before.

        Args:
            post_id: Post ID to check

        Returns:
            True if post was seen before
        """
        return post_id in self.seen_posts

    def mark_post_seen(self, post_id: str):
        """Mark a post as seen.

        Args:
            post_id: Post ID to mark
        """
        self.seen_posts.add(post_id)
        self._save_seen_posts()

    def add_post(
        self,
        post_data: Dict,
        profile_name: str,
        profile_url: str,
        comment_data: Dict,
        comment_style: str,
        posted: bool = False
    ):
        """Add a post and its generated comment to the Excel file.

        Args:
            post_data: Dictionary with post information
            profile_name: Name of the profile
            profile_url: URL of the profile
            comment_data: Dictionary with comment, tone, and reasoning
            comment_style: Style used for comment generation
            posted: Whether the comment was actually posted
        """
        print(f"{Fore.CYAN}Adding post to Excel tracker...{Style.RESET_ALL}")

        # Read existing data
        try:
            df = pd.read_excel(self.excel_file, sheet_name='Posts')
        except Exception:
            df = pd.DataFrame()

        # Create new row
        new_row = {
            'Post ID': post_data.get('post_id', ''),
            'Profile Name': profile_name,
            'Profile URL': profile_url,
            'Post URL': post_data.get('url', ''),
            'Post Date': post_data.get('timestamp', datetime.now()).strftime('%Y-%m-%d %H:%M:%S'),
            'Scraped At': post_data.get('scraped_at', datetime.now()).strftime('%Y-%m-%d %H:%M:%S'),
            'Post Content': post_data.get('content', '')[:1000],  # Limit length
            'Detected Tone': comment_data.get('tone', 'unknown'),
            'Generated Comment': comment_data.get('comment', ''),
            'Comment Style': comment_style,
            'Reasoning': comment_data.get('reasoning', ''),
            'Likes': post_data.get('likes', 0),
            'Comments': post_data.get('comments', 0),
            'Shares': post_data.get('shares', 0),
            'Posted Comment': 'Yes' if posted else 'No',
            'Posted At': datetime.now().strftime('%Y-%m-%d %H:%M:%S') if posted else '',
            'Notes': ''
        }

        # Append row
        df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)

        # Save to Excel
        df.to_excel(self.excel_file, index=False, sheet_name='Posts')

        # Reapply formatting
        self._format_excel()

        # Mark post as seen
        self.mark_post_seen(post_data.get('post_id', ''))

        print(f"{Fore.GREEN}Post added to Excel tracker{Style.RESET_ALL}")

    def get_stats(self) -> Dict:
        """Get tracking statistics.

        Returns:
            Dictionary with statistics
        """
        try:
            df = pd.read_excel(self.excel_file, sheet_name='Posts')

            stats = {
                'total_posts': len(df),
                'posts_today': len(df[pd.to_datetime(df['Scraped At']).dt.date == datetime.now().date()]),
                'comments_generated': len(df[df['Generated Comment'].notna()]),
                'comments_posted': len(df[df['Posted Comment'] == 'Yes']),
                'unique_profiles': df['Profile Name'].nunique(),
                'avg_likes': df['Likes'].mean() if 'Likes' in df.columns else 0,
            }

            return stats

        except Exception:
            return {
                'total_posts': 0,
                'posts_today': 0,
                'comments_generated': 0,
                'comments_posted': 0,
                'unique_profiles': 0,
                'avg_likes': 0,
            }

    def print_stats(self):
        """Print tracking statistics to console."""
        stats = self.get_stats()

        print(f"\n{Fore.CYAN}{'='*50}")
        print(f"Tracking Statistics")
        print(f"{'='*50}{Style.RESET_ALL}")
        print(f"Total Posts Tracked: {stats['total_posts']}")
        print(f"Posts Today: {stats['posts_today']}")
        print(f"Comments Generated: {stats['comments_generated']}")
        print(f"Comments Posted: {stats['comments_posted']}")
        print(f"Unique Profiles: {stats['unique_profiles']}")
        print(f"Avg Likes per Post: {stats['avg_likes']:.1f}")
        print(f"{Fore.CYAN}{'='*50}{Style.RESET_ALL}\n")
